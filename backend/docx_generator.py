"""
ATS-friendly DOCX Resume Generator.
Generates clean, single-column Microsoft Word documents adhering to strict ATS readability guidelines.
No floating text boxes, multi-column tables, or complex decorative shapes.
"""

import io
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_bottom_border(paragraph):
    """Add a subtle clean bottom border to section headings."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')  # 0.75 pt
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '0F172A')
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_resume_docx(data: Dict[str, Any]) -> io.BytesIO:
    """Generate ATS-friendly DOCX file from structured resume dictionary."""
    doc = Document()

    # Set page margins to standard 0.45-0.5 inch for strict 1-page fit
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)

    # 1. Header: Candidate Name
    candidate_name = data.get("candidate_name", "Candidate Name").strip()
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(1)
    name_run = name_p.add_run(candidate_name)
    name_run.font.size = Pt(18)
    name_run.font.bold = True
    name_run.font.color.rgb = RGBColor(15, 23, 42)

    # Target Role Subtitle
    target_role = data.get("target_role", "").strip()
    if target_role:
        role_p = doc.add_paragraph()
        role_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        role_p.paragraph_format.space_before = Pt(0)
        role_p.paragraph_format.space_after = Pt(2)
        role_run = role_p.add_run(target_role)
        role_run.font.size = Pt(10)
        role_run.font.bold = True
        role_run.font.color.rgb = RGBColor(37, 99, 235)  # Cobalt accent

    # Contact line
    contact_parts = []
    contact_info = data.get("contact_info", {})
    if isinstance(contact_info, dict):
        if contact_info.get("email"): contact_parts.append(contact_info["email"])
        if contact_info.get("phone"): contact_parts.append(contact_info["phone"])
        if contact_info.get("location"): contact_parts.append(contact_info["location"])
        if contact_info.get("linkedin"): contact_parts.append(contact_info["linkedin"])
        if contact_info.get("github"): contact_parts.append(contact_info["github"])
        if contact_info.get("portfolio"): contact_parts.append(contact_info["portfolio"])
    elif isinstance(contact_info, list):
        contact_parts = [str(c) for c in contact_info if c]

    if contact_parts:
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_p.paragraph_format.space_before = Pt(0)
        contact_p.paragraph_format.space_after = Pt(6)
        contact_run = contact_p.add_run("  |  ".join(contact_parts))
        contact_run.font.size = Pt(9)
        contact_run.font.color.rgb = RGBColor(71, 85, 105)

    def add_section_header(title: str):
        sec_p = doc.add_paragraph()
        sec_p.paragraph_format.space_before = Pt(7)
        sec_p.paragraph_format.space_after = Pt(2)
        sec_p.paragraph_format.keep_with_next = True
        run = sec_p.add_run(title.upper())
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        add_bottom_border(sec_p)

    # 2. Professional Summary
    summary = data.get("professional_summary", "").strip()
    if summary:
        add_section_header("Professional Summary")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(summary)
        run.font.size = Pt(9.5)

    # 3. Core Competencies & Skills
    skills = data.get("skills", {})
    if skills:
        add_section_header("Skills & Competencies")
        if isinstance(skills, dict):
            for category, skill_list in skills.items():
                if skill_list:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    cat_run = p.add_run(f"{category}: ")
                    cat_run.bold = True
                    cat_run.font.size = Pt(9.5)
                    
                    skills_str = ", ".join(skill_list) if isinstance(skill_list, list) else str(skill_list)
                    s_run = p.add_run(skills_str)
                    s_run.font.size = Pt(9.5)
        elif isinstance(skills, list):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.add_run(", ".join(skills))

    # 4. Work Experience
    experience = data.get("work_experience", [])
    if experience:
        add_section_header("Professional Experience")
        for exp in experience:
            # Job Title and Company Header line
            title_p = doc.add_paragraph()
            title_p.paragraph_format.space_before = Pt(4)
            title_p.paragraph_format.space_after = Pt(1)
            title_p.paragraph_format.keep_with_next = True

            role_run = title_p.add_run(exp.get("role", "Role"))
            role_run.bold = True
            role_run.font.size = Pt(10)

            comp = exp.get("company", "")
            if comp:
                c_run = title_p.add_run(f" | {comp}")
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(71, 85, 105)

            dates = exp.get("dates", "")
            loc = exp.get("location", "")
            meta_parts = [p for p in [loc, dates] if p]
            if meta_parts:
                meta_p = doc.add_paragraph()
                meta_p.paragraph_format.space_before = Pt(0)
                meta_p.paragraph_format.space_after = Pt(2)
                meta_p.paragraph_format.keep_with_next = True
                m_run = meta_p.add_run(" - ".join(meta_parts))
                m_run.font.size = Pt(8.5)
                m_run.font.italic = True
                m_run.font.color.rgb = RGBColor(100, 116, 139)

            # Bullet points
            bullets = exp.get("bullets", [])
            for bullet in bullets:
                if bullet.strip():
                    bp = doc.add_paragraph(style='List Bullet')
                    bp.paragraph_format.space_before = Pt(0.5)
                    bp.paragraph_format.space_after = Pt(1.5)
                    bp.paragraph_format.line_spacing = 1.12
                    b_run = bp.add_run(bullet.strip())
                    b_run.font.size = Pt(9)

    # 5. Projects
    projects = data.get("projects", [])
    if projects:
        add_section_header("Key Technical Projects")
        for proj in projects:
            p_head = doc.add_paragraph()
            p_head.paragraph_format.space_before = Pt(4)
            p_head.paragraph_format.space_after = Pt(1)
            p_head.paragraph_format.keep_with_next = True

            p_name = p_head.add_run(proj.get("name", "Project"))
            p_name.bold = True
            p_name.font.size = Pt(9.5)

            tech = proj.get("technologies", [])
            tech_str = ", ".join(tech) if isinstance(tech, list) else str(tech)
            if tech_str:
                t_run = p_head.add_run(f" ({tech_str})")
                t_run.font.size = Pt(9)
                t_run.font.color.rgb = RGBColor(71, 85, 105)

            bullets = proj.get("bullets", [])
            for b in bullets:
                if b.strip():
                    bp = doc.add_paragraph(style='List Bullet')
                    bp.paragraph_format.space_before = Pt(0.5)
                    bp.paragraph_format.space_after = Pt(1.5)
                    bp.paragraph_format.line_spacing = 1.12
                    b_run = bp.add_run(b.strip())
                    b_run.font.size = Pt(9)

    # 6. Education
    education = data.get("education", [])
    if education:
        add_section_header("Education")
        for edu in education:
            edu_p = doc.add_paragraph()
            edu_p.paragraph_format.space_before = Pt(4)
            edu_p.paragraph_format.space_after = Pt(2)
            deg = edu_p.add_run(edu.get("degree", "Degree"))
            deg.bold = True
            deg.font.size = Pt(10)

            inst = edu.get("institution", "")
            if inst:
                inst_run = edu_p.add_run(f", {inst}")
                inst_run.font.size = Pt(10)

            dates = edu.get("graduation_year", "") or edu.get("dates", "")
            if dates:
                d_run = edu_p.add_run(f" ({dates})")
                d_run.font.size = Pt(9.5)
                d_run.font.color.rgb = RGBColor(100, 116, 139)

    # 7. Certifications (strictly only if real, verified, and not N/A)
    certs = data.get("certifications", [])
    valid_certs = []
    for cert in certs:
        if not isinstance(cert, dict):
            continue
        c_name = str(cert.get("name", "")).strip()
        c_issuer = str(cert.get("issuer", "")).strip()
        if not c_name or c_name.lower() in ["n/a", "none", "not applicable"]:
            continue
        if "n/a" in c_name.lower() or "none" in c_name.lower() or "(n/a)" in c_name.lower():
            continue
        if c_issuer.lower() in ["n/a", "none"] and ("certification" in c_name.lower() or "certificate" in c_name.lower()):
            continue
        valid_certs.append(cert)

    if valid_certs:
        add_section_header("Certifications & Credentials")
        for cert in valid_certs:
            cert_p = doc.add_paragraph(style='List Bullet')
            cert_p.paragraph_format.space_before = Pt(1)
            cert_p.paragraph_format.space_after = Pt(2)
            c_name = cert.get("name", "")
            run = cert_p.add_run(c_name)
            run.font.size = Pt(9.5)
            issuer = cert.get("issuer", "")
            if issuer and "n/a" not in str(issuer).lower():
                i_run = cert_p.add_run(f" — {issuer}")
                i_run.font.size = Pt(9)
                i_run.font.color.rgb = RGBColor(100, 116, 139)
            year = cert.get("year", "")
            if year and "n/a" not in str(year).lower() and str(year).lower() != "year":
                y_run = cert_p.add_run(f" ({year})")
                y_run.font.size = Pt(8.5)
                y_run.font.color.rgb = RGBColor(148, 163, 184)

    # Save to memory buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
